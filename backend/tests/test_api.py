"""End-to-end synthetic acceptance tests; no industrial data in the repository."""

import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from procurement_app.api.main import create_app
from procurement_app.config import Settings
from procurement_app.data.store import init_store
from procurement_app.services.context import AppContext, set_context


@pytest.fixture
def client(seed_source, tmp_path):
    ctx = AppContext(
        Settings(mode="demo", data_source="local", as_of=dt.date(2026, 9, 21), horizon_days=30),
        seed_source,
        init_store(f"sqlite:///{tmp_path}/test.db"),
    )
    set_context(ctx)
    c = TestClient(create_app(), headers={"x-forwarded-email": "planner@example.com", "x-procurement-request": "1"})

    def update_version(response):
        if response.is_success and "x-data-revision" in response.headers:
            c.headers["if-match"] = response.headers["x-data-revision"]

    c.event_hooks["response"] = [update_version]
    c.get("/api/config")
    yield c
    c.close()
    set_context(None)


def projection(c, scenario=None):
    r = c.get("/api/articles/DEMO-001/projection", params={"scenario_id": scenario} if scenario else {})
    assert r.status_code == 200, r.text
    return r.json()


def flows(r, key):
    return next(x["values"] for x in r["series"] if x["key"] == key)


def test_all_read_pages_and_exports(client):
    for url in [
        "/api/config",
        "/api/cockpit",
        "/api/reference/articles",
        "/api/reference/suppliers",
        "/api/reference/links",
        "/api/reference/bom",
        "/api/reference/programs",
        "/api/reference/plan",
        "/api/params/effective",
        "/api/params/schema",
        "/api/params/overrides",
        "/api/entries/orders",
        "/api/entries/receipts",
        "/api/entries/production",
        "/api/entries/adjustments",
        "/api/entries/cells",
        "/api/scenarios",
        "/api/pdp/versions",
        "/api/audit",
        "/api/exports/alerts.xlsx",
        "/api/exports/orders.xlsx",
    ]:
        r = client.get(url)
        assert r.status_code == 200, (url, r.text)
    assert client.get("/api/cockpit").json()["kpis"]["articles"] == 8
    day = projection(client)
    week = client.get("/api/articles/DEMO-001/projection?granularity=week").json()
    assert sum(flows(day, "demand")) == pytest.approx(sum(flows(week, "demand")))


def test_partial_receipt_and_reversal(client):
    r = client.post(
        "/api/entries/orders",
        json={"article_id": "DEMO-001", "supplier_id": "SUP-2", "expected_date": "2026-09-25", "qty": 100},
    )
    assert r.status_code == 201, r.text
    oid = r.json()["id"]
    r = client.post(
        "/api/entries/receipts",
        json={"article_id": "DEMO-001", "order_id": oid, "receipt_date": "2026-09-24", "qty": 60},
    )
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    p = projection(client)
    i = p["periods"].index("2026-09-25")
    assert flows(p, "supply_firm")[i] == 40
    assert (
        client.post(
            "/api/entries/receipts",
            json={"article_id": "DEMO-001", "order_id": oid, "receipt_date": "2026-09-24", "qty": 50},
        ).status_code
        == 422
    )
    assert client.delete(f"/api/entries/receipts/{rid}").status_code == 204
    assert flows(projection(client), "supply_firm")[i] == 100
    assert any(x["action"] == "snapshot_create" and x["entity_id"] for x in client.get("/api/audit").json())


def test_invalid_inputs_and_concurrency(client):
    stale = client.headers["if-match"]
    r = client.put(
        "/api/entries/cells", json={"article_id": "DEMO-001", "date": "2026-09-25", "expression": "2*600-50"}
    )
    assert r.status_code == 200, r.text
    assert r.json()["qty"] == 1150
    assert (
        client.put(
            "/api/entries/cells",
            headers={"if-match": stale},
            json={"article_id": "DEMO-001", "date": "2026-09-26", "expression": "1"},
        ).status_code
        == 409
    )
    assert (
        client.put(
            "/api/entries/cells",
            json={"article_id": "DEMO-001", "date": "2026-09-25", "expression": '__import__("os")'},
        ).status_code
        == 422
    )
    assert client.get("/api/cockpit?production_mode=nonsense").status_code == 422
    assert (
        client.post(
            "/api/entries/orders",
            json={"article_id": "DEMO-001", "supplier_id": "NOPE", "expected_date": "2026-09-25", "qty": 100},
        ).status_code
        == 422
    )
    assert (
        client.put(
            "/api/params/overrides", json={"scope": "global", "field": "horizon_days", "value": 999999}
        ).status_code
        == 422
    )
    assert client.get("/api/cockpit").status_code == 200


def test_frozen_scenario_and_scoped_cbn(client):
    r = client.post(
        "/api/scenarios", json={"name": "PDP +20%", "events": [{"kind": "plan_factor", "payload": {"factor": 1.2}}]}
    )
    assert r.status_code == 201, r.text
    sc = r.json()["id"]
    before = projection(client, sc)
    assert (
        client.post(
            "/api/entries/adjustments", json={"article_id": "DEMO-001", "date": "2026-09-24", "qty": 100}
        ).status_code
        == 201
    )
    assert flows(projection(client, sc), "stock_sim") == flows(before, "stock_sim")
    base = flows(projection(client), "stock_sim")
    r = client.post(
        "/api/cbn/run", json={"scenario_id": sc, "article_ids": ["DEMO-001"], "params": {"horizon_days": 30}}
    )
    assert r.status_code == 200, r.text
    assert r.json()["proposals"] > 0
    assert flows(projection(client), "stock_sim") == base
    assert client.get("/api/entries/cells").json() == []
    first = r.json()
    r = client.post(
        "/api/cbn/run", json={"scenario_id": sc, "article_ids": ["DEMO-001"], "params": {"horizon_days": 30}}
    )
    assert r.status_code == 200, r.text
    assert r.json()["qty"] == first["qty"]
    assert client.get(f"/api/scenarios/{sc}/compare?horizon_days=30").status_code == 200


def test_excel_roundtrip_all_flows_and_idempotency(client):
    res = client.get("/api/exports/simulation.xlsx?article_ids=DEMO-001&granularity=week&horizon_days=30")
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    assert "SAISIES" not in wb and "CARNET_COMMANDES" not in wb
    assert {"SIMULATION", "HEBDOMADAIRE", "LIRE_MOI", "_FORMAT"} == set(wb.sheetnames)
    assert client.post("/api/imports/entries", files={"file": ("same.xlsx", res.content)}).json()["created"] == 0
    ws = wb["SIMULATION"]
    ws["E8"] = "=2*600-50"
    ws["F10"] = 17  # simulated order, signed adjustment
    stream = io.BytesIO()
    wb.save(stream)
    content = stream.getvalue()
    preview = client.post("/api/imports/preview", files={"file": ("edit.xlsx", content)})
    assert preview.status_code == 200, preview.text
    assert preview.json()["count"] == 2
    imported = client.post("/api/imports/entries", files={"file": ("edit.xlsx", content)})
    assert imported.status_code == 201, imported.text
    assert imported.json()["created"] == 2
    sc = client.get("/api/scenarios").json()[0]["id"]
    p = projection(client, sc)
    assert flows(p, "supply_planned")[p["periods"].index("2026-09-21")] == 1150
    assert flows(p, "adjustments")[p["periods"].index("2026-09-22")] == 17
    assert client.post("/api/imports/entries", files={"file": ("edit.xlsx", content)}).json()["created"] == 0
    assert client.get("/api/entries/orders").json() == []


def test_pdp_partial_week_override_and_validation(client):
    from openpyxl import Workbook

    w = Workbook()
    w.active.append(["program_id", "week_start", "qty"])
    w.active.append(["VEH-A", dt.date(2026, 9, 21), 100])
    buf = io.BytesIO()
    w.save(buf)
    r = client.post("/api/pdp/import", files={"file": ("pdp.xlsx", buf.getvalue())})
    assert r.status_code == 201, r.text
    p = projection(client)
    assert flows(p, "demand")[p["periods"].index("2026-09-28")] > 0
    assert client.post(f"/api/pdp/versions/{r.json()['version']['id']}/deactivate").status_code == 200


def test_accept_ignore_and_duplicate_decision(client):
    r = client.post("/api/cbn/run", json={"article_ids": ["DEMO-001"], "params": {"horizon_days": 30}})
    assert r.status_code == 200, r.text
    p = r.json()["items"][0]
    before = projection(client)
    r = client.post("/api/cbn/decisions/" + p["proposal_id"], json={"action": "accept", "qty": p["qty"]})
    assert r.status_code == 200, r.text
    assert flows(projection(client), "supply_planned") == flows(before, "supply_planned")
    assert client.post("/api/cbn/decisions/" + p["proposal_id"], json={"action": "accept"}).status_code == 409
    assert client.get("/api/entries/orders").json()[0]["order_type"] == "PLANNED"


def test_private_scenarios_reader_and_admin_roles(client):
    from procurement_app.services.context import get_context

    ctx = get_context()
    ctx.settings.mode = "production"
    ctx.settings.editors = "planner@example.com,other@example.com"
    ctx.settings.admins = "admin@example.com"
    sc = client.post("/api/scenarios", json={"name": "Private"}).json()["id"]
    r = client.get("/api/scenarios", headers={"x-forwarded-email": "other@example.com"})
    assert r.json() == []
    assert client.get(f"/api/scenarios/{sc}", headers={"x-forwarded-email": "other@example.com"}).status_code == 403
    assert (
        client.post(
            "/api/cbn/run", headers={"x-forwarded-email": "other@example.com"}, json={"scenario_id": sc}
        ).status_code
        == 403
    )
    assert (
        client.post(
            "/api/entries/adjustments",
            headers={"x-forwarded-email": "reader@example.com"},
            json={"article_id": "DEMO-001", "date": "2026-09-24", "qty": 1},
        ).status_code
        == 403
    )
    assert (
        client.put("/api/params/overrides", json={"scope": "global", "field": "horizon_days", "value": 30}).status_code
        == 403
    )
    assert client.get("/api/config", headers={"x-forwarded-email": ""}).status_code == 401


def test_scenario_clone_and_v1_import(client):
    from procurement.demo import dataset
    from procurement_app.services.migration import from_v1

    original = dataset()
    converted, params = from_v1(original)
    assert len(converted.articles) == len(original.items)
    r = client.post("/api/imports/v1", files={"file": ("v1.json", original.model_dump_json().encode())})
    assert r.status_code == 201, r.text
    sc = client.get("/api/scenarios").json()[0]["id"]
    cloned = client.post(f"/api/scenarios/{sc}/clone")
    assert cloned.status_code == 200, cloned.text
    assert cloned.json()["id"] != sc
