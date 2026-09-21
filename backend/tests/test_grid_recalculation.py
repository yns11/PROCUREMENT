"""Recalculate with LibreOffice, then compare every stock/shortage/target/coverage cell."""

import io
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from procurement_app.engine import run_mrp
from procurement_app.engine.models import EngineParams, SimCell
from procurement_app.services.grid_excel import export_grid, parse_grid

from .test_engine_units import MON, make_dataset


@pytest.mark.parametrize("policy", ["backlog", "lost"])
@pytest.mark.parametrize("unit", ["calendar", "working"])
def test_native_excel_matches_engine(tmp_path, policy, unit):
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        pytest.skip("LibreOffice absent; run scripts/check_excel.sh or CI Excel job")
    ds = make_dataset()
    ds.holidays = [MON.replace(day=23)]
    params = EngineParams(as_of=MON, horizon_days=20, coverage_unit=unit, shortage_policy=policy)
    original = run_mrp(ds, params)
    content, manifest = export_grid(original, granularity="week", meta={"holidays": ds.holidays}, export_id="test")
    wb = load_workbook(io.BytesIO(content))
    ws = wb["SIMULATION"]
    ws["E8"] = "=2*600-50"
    ws["F10"] = -25
    buf = io.BytesIO()
    wb.save(buf)
    for c in parse_grid(buf.getvalue(), manifest):
        import datetime as dt

        ds.cells.append(SimCell(c["article_id"], dt.date.fromisoformat(c["date"]), c["kind"], c["qty"]))
    expected = run_mrp(ds, params).articles["A1"]
    source = tmp_path / "input.xlsx"
    source.write_bytes(buf.getvalue())
    out = tmp_path / "calculated"
    out.mkdir()
    p = subprocess.run(
        [
            office,
            "-env:UserInstallation=file://" + str(tmp_path / "profile"),
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert p.returncode == 0, p.stdout + p.stderr
    calc = load_workbook(out / "input.xlsx", data_only=True)["SIMULATION"]
    rows = manifest["rows"]["A1"]
    for layer in ("firm", "forecast", "sim"):
        for prefix, series in [
            ("net", "stock_" + layer + "_net"),
            ("physical", "stock_" + layer),
            ("shortage", "shortage_" + layer),
        ]:
            for j, d in enumerate(manifest["dates"], 5):
                i = next(i for i, x in enumerate(expected.dates) if x.isoformat() == d)
                assert calc.cell(rows[prefix + "_" + layer], j).value == pytest.approx(getattr(expected, series)[i]), (
                    policy,
                    unit,
                    prefix,
                    layer,
                    d,
                )
    for key, series in [("target", "target_stock"), ("coverage", "coverage_sim")]:
        for j, d in enumerate(manifest["dates"], 5):
            i = next(i for i, x in enumerate(expected.dates) if x.isoformat() == d)
            assert calc.cell(rows[key], j).value == pytest.approx(getattr(expected, series)[i]), (key, d)
