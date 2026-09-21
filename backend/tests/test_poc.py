"""Executable management presentation scenarios on entirely synthetic inputs."""

import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from procurement_app.api.main import create_app
from procurement_app.config import Settings
from procurement_app.data.poc import DAY, PocSource
from procurement_app.data.store import init_store
from procurement_app.services.context import AppContext, set_context


@pytest.fixture
def poc_client(tmp_path):
    factory = init_store(f"sqlite:///{tmp_path}/poc.db")
    ctx = AppContext(
        Settings(mode="demo", poc=True, data_source="local", as_of=DAY, horizon_days=60, history_days=0),
        PocSource(factory),
        factory,
    )
    set_context(ctx)
    c = TestClient(create_app(), headers={"x-procurement-request": "1"})

    def version(r):
        if r.is_success and "x-data-revision" in r.headers:
            c.headers["if-match"] = r.headers["x-data-revision"]

    c.event_hooks["response"] = [version]
    assert c.get("/api/config").status_code == 200
    yield c
    c.close()
    set_context(None)


def project(c):
    r = c.get("/api/articles/DEMO-001/projection", params={"horizon_days": 60, "granularity": "day"})
    assert r.status_code == 200, r.text
    return r.json()


def value(p, series, date):
    values = next(s["values"] for s in p["series"] if s["key"] == series)
    return values[p["periods"].index(date)]


def table(c, name):
    r = c.get(f"/api/poc/tables/{name}")
    assert r.status_code == 200, r.text
    return r.json()["rows"]


def save(c, name, rows):
    r = c.put(f"/api/poc/tables/{name}", json={"rows": rows})
    assert r.status_code == 200, r.text
    return r.json()["rows"]


def cbn(c):
    r = c.post("/api/cbn/run", json={"article_ids": ["DEMO-001"], "params": {"horizon_days": 60}})
    assert r.status_code == 200, r.text
    return r.json()


def test_poc_mondays_and_repeatable_cbn(poc_client):
    c = poc_client
    rows = table(c, "orders")
    assert rows[0]["expected_date"] == "2026-09-24"  # firm stays Thursday
    assert rows[1]["expected_date"] == "2026-09-28"
    p = project(c)
    assert value(p, "supply_forecast", "2026-09-28") == 200
    assert value(p, "supply_forecast", "2026-09-30") == 0
    first = cbn(c)
    assert first["items"]
    assert all(dt.date.fromisoformat(x["delivery_date"]).weekday() == 0 for x in first["items"])
    assert cbn(c)["qty"] == first["qty"]


def test_pdp_increase_changes_stock_and_cbn(poc_client):
    c = poc_client
    before = project(c)
    first = cbn(c)
    rows = table(c, "pdp")
    rows[1]["qty"] = 400
    save(c, "pdp", rows)
    after = project(c)
    assert value(after, "demand", "2026-09-28") == 160
    assert value(after, "stock_firm", "2026-09-28") == value(before, "stock_firm", "2026-09-28") - 60
    second = cbn(c)
    assert second["qty"] > first["qty"]
    print(
        "PDP scenario",
        {
            "cbn_before": first["qty"],
            "cbn_after": second["qty"],
            "stock_before": value(before, "stock_firm", "2026-09-28"),
            "stock_after": value(after, "stock_firm", "2026-09-28"),
        },
    )


@pytest.mark.parametrize("source", ["manual", "cbn"])
def test_firm_absorbs_simulation_once_and_reverses(poc_client, source):
    c = poc_client
    if source == "manual":
        day = "2026-09-28"
        assert (
            c.put("/api/entries/cells", json={"article_id": "DEMO-001", "date": day, "expression": "500"}).status_code
            == 200
        )
    else:
        result = cbn(c)
        day = result["items"][0]["delivery_date"]
    before = project(c)
    qty = value(before, "supply_planned", day)
    assert qty > 0
    initial = table(c, "orders")
    row = {
        "order_id": "CONFIRM-001",
        "article_id": "DEMO-001",
        "supplier_id": "FOURN-01",
        "expected_date": day,
        "qty": qty,
        "order_type": "FIRM",
    }
    save(c, "orders", initial + [row])
    exported = c.get("/api/exports/simulation.xlsx?article_ids=DEMO-001&horizon_days=60")
    assert exported.status_code == 200, exported.text
    grid = load_workbook(io.BytesIO(exported.content))["SIMULATION"]
    column = next(i for i in range(5, grid.max_column + 1) if str(grid.cell(2, i).value)[:10] == day)
    simulated_row = next(r[0].row for r in grid if r[2].value == "Commandes simulées")
    assert grid.cell(simulated_row, column).value == 0
    for _ in range(2):
        after = project(c)
        assert value(after, "supply_planned", day) == 0
        assert value(after, "supply_firm", day) == value(before, "supply_firm", day) + qty
        assert value(after, "stock_sim", day) == value(before, "stock_sim", day)
    row["qty"] = qty / 2
    save(c, "orders", initial + [row])
    assert value(project(c), "supply_planned", day) == qty / 2
    save(c, "orders", initial)
    assert value(project(c), "supply_planned", day) == qty


def test_reference_changes_and_closed_monday(poc_client):
    from procurement_app.services.context import get_context

    c = poc_client
    rows = table(c, "bom")
    rows[0]["qty_per"] = 3
    save(c, "bom", rows)
    assert value(project(c), "demand", "2026-09-28") == 150
    ctx = get_context()
    ctx.settings.holidays = "2026-09-28"
    ctx.bump()
    result = cbn(c)
    assert result["items"]
    assert all(x["delivery_date"] >= "2026-10-05" for x in result["items"])
    assert all(dt.date.fromisoformat(x["delivery_date"]).weekday() == 0 for x in result["items"])


def test_inputs_roundtrip_atomic_validation_and_reset(poc_client):
    c = poc_client
    for name in ("articles", "bom", "pdp", "orders"):
        original = table(c, name)
        f = c.get(f"/api/poc/templates/{name}.csv")
        revision = c.headers["if-match"]
        r = c.post(f"/api/poc/parse/{name}", files={"file": ("example.csv", f.content)})
        assert r.status_code == 200, r.text
        assert c.headers["if-match"] == revision  # parsing is read-only
        assert r.json()["rows"] == original
        assert save(c, name, r.json()["rows"]) == original
    rows = table(c, "bom")
    rows[0]["article_id"] = "UNKNOWN"
    assert c.put("/api/poc/tables/bom", json={"rows": rows}).status_code == 422
    assert table(c, "bom")[0]["article_id"] == "DEMO-001"
    rows = table(c, "orders")
    rows[1]["expected_date"] = "2027-01-03"
    assert save(c, "orders", rows)[1]["expected_date"] == "2026-12-28"
    assert c.post("/api/poc/reset").status_code == 200
    assert table(c, "orders")[1]["expected_date"] == "2026-09-28"
    wb = Workbook()
    wb.active.append(["program_id", "week_start", "qty"])
    wb.active.append(["MOTEUR-A", dt.datetime(2026, 9, 28), 400])
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/api/poc/parse/pdp", files={"file": ("pdp.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    assert r.json()["rows"][0]["qty"] == 400
    wb.active["C2"] = "=2*200"
    buf = io.BytesIO()
    wb.save(buf)
    assert c.post("/api/poc/parse/pdp", files={"file": ("bad.xlsx", buf.getvalue())}).status_code == 422
    assert c.post("/api/poc/parse/pdp", files={"file": ("broken.xlsx", b"invalid")}).status_code == 422
