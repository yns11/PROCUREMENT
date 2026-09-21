"""ERP text must remain literal even when a reference starts with '='."""

import io
from dataclasses import replace

from openpyxl import load_workbook

from procurement_app.engine import run_mrp
from procurement_app.engine.models import EngineParams, OrderLine, OrderType
from procurement_app.services.excel_service import alerts_workbook, orders_workbook
from procurement_app.services.grid_excel import export_grid

from .test_engine_units import MON, make_dataset


def test_untrusted_text_is_not_exported_as_formulas():
    text = '=HYPERLINK("https://example.invalid","ERP")'
    ds = make_dataset()
    ds.articles = [replace(ds.articles[0], article_id=text, designation=text, unit=text)]
    ds.links = [replace(x, article_id=text) for x in ds.links]
    ds.bom = [replace(x, article_id=text) for x in ds.bom]
    ds.stock = [replace(x, article_id=text, qty_on_hand=0) for x in ds.stock]
    ds.orders = [OrderLine(text, text, "S1", MON, 1, order_type=OrderType.FIRM)]
    result = run_mrp(ds, EngineParams(as_of=MON, horizon_days=10, generate_proposals=False))
    grid, manifest = export_grid(result, granularity="week")
    wb = load_workbook(io.BytesIO(grid))
    row = manifest["rows"][text]["demand"]
    for cell in [
        wb["SIMULATION"].cell(row, 1),
        wb["SIMULATION"].cell(row, 2),
        wb["HEBDOMADAIRE"]["A2"],
        wb["HEBDOMADAIRE"]["C2"],
    ]:
        assert cell.value.startswith("=")
        assert cell.data_type == "s"
    for content in [orders_workbook(result), alerts_workbook(result)]:
        ws = load_workbook(io.BytesIO(content)).active
        assert ws.max_row > 1
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s"
