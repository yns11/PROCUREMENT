"""PDP parsing and standalone extracts; simulation grid is in grid_excel.py."""

from __future__ import annotations

import datetime as dt
import io
import re
from dataclasses import dataclass
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..engine.calendar import iso_week_label, iso_week_monday
from ..engine.models import MrpResult
from .grid_excel import export_grid

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="EEF1F6")
INPUT_FILL = PatternFill("solid", fgColor="FFFBEA")
HELPER_FONT = Font(color="98A2B3", italic=True)
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
YELLOW_FILL = PatternFill("solid", fgColor="FFF3CD")
GREEN_FILL = PatternFill("solid", fgColor="D1E7DD")
BLUE_FONT = Font(color="255D95", bold=True)  # Excel convention: blue = input, black = formula
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D0D5DD")
DATE_FMT = "yyyy-mm-dd"
QTY_FMT = "#,##0.###"

# Fixed cells of the PARAMETRES sheet referenced by the formulas
P_SHORTAGE = "PARAMETRES!$B$5"
P_TARGET = "PARAMETRES!$B$6"
P_TIE = "PARAMETRES!$B$7"

SPARE_ROWS = 300  # blank rows kept in the SUMIFS ranges for planner additions
ENTRY_ROWS = 2000  # rows scanned in SAISIES
ENTRY_FORMAT_ROWS = 200  # rows of SAISIES pre-formatted as dates (Excel recognises typed dates anyway)

# Row layout of one article block in SIMULATION (offset → key, label, kind)
BLOCK = [
    ("demand", "Besoin", "input"),
    ("orders_firm", "Commandes fermes (carnet)", "formula"),
    ("orders_forecast", "Commandes prévisionnelles ERP (carnet)", "formula"),
    ("sim_orders", "Commandes simulées", "input"),
    ("entries_orders", "Saisies : commandes (SAISIES)", "formula"),
    ("entries_receipts", "Saisies : réceptions & ajustements (SAISIES, carnet)", "formula"),
    ("known_receipts", "Réceptions & ajustements connus", "input"),
    ("stock_firm", "Stock ferme", "stock"),
    ("stock_forecast", "Stock prévisionnel", "stock"),
    ("stock_sim", "Stock simulé", "stock"),
    ("shortage_sim", "Manque simulé (besoin non servi)", "formula"),
    ("target", "Stock cible", "formula"),
    ("coverage", "Couverture simulée (périodes)", "formula"),
    ("cum_demand", "Besoin cumulé (aide au calcul)", "helper"),
]
BLOCK_ROWS = len(BLOCK)
ROW = {key: i for i, (key, _, _) in enumerate(BLOCK)}
SIM_HEADER_ROWS = 3  # label / period start / period end
SIM_FIRST_COL = 5  # E


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: Iterable[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def periods(dates: list[dt.date], granularity: str, start: dt.date) -> list[tuple[str, list[int]]]:
    """Group the day indexes of ``dates`` (from ``start``) into labelled periods."""
    out: dict[str, list[int]] = {}
    for i, d in enumerate(dates):
        if d < start:
            continue
        label = d.isoformat() if granularity == "day" else iso_week_label(d)
        out.setdefault(label, []).append(i)
    return list(out.items())


def _date_cell(ws, row: int, col: int, value: dt.date | None):
    cell = ws.cell(row, col, value)
    cell.number_format = DATE_FMT
    return cell


# =============================================================================
# Simulation workbook
# =============================================================================
def _alerts_sheet(ws, result: MrpResult, ids: list[str]) -> None:
    headers = ["Article", "Désignation", "Type", "Sévérité", "Périmètre", "Date", "Valeur", "Message"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    _style_header(ws, 1, len(headers))
    r = 2
    for aid in ids:
        ar = result.articles.get(aid)
        if not ar:
            continue
        for a in ar.alerts:
            ws.append(
                [
                    aid,
                    ar.article.designation,
                    a.alert_type.value,
                    a.severity.value,
                    a.scope,
                    a.date.isoformat() if a.date else "",
                    a.value,
                    a.message,
                ]
            )
            fill = (
                RED_FILL if a.severity.value == "critical" else YELLOW_FILL if a.severity.value == "warning" else None
            )
            if fill:
                ws.cell(r, 4).fill = fill
            r += 1
    _widths(ws, (13, 28, 18, 10, 10, 12, 12, 90))
    ws.freeze_panes = "A2"


ORDER_HEADERS = [
    "Article",
    "Désignation",
    "Type",
    "Référence",
    "Fournisseur",
    "Date attendue",
    "Quantité",
    "Origine",
    "Retard",
    "Reçu (saisie)",
    "Date réception (saisie)",
    "Statut (saisie : RECUE / ANNULEE)",
    "Reste à livrer",
]


def _orders_sheet(ws, result: MrpResult, ids: list[str], start: dt.date | None = None, live: bool = False) -> int:
    """Open order book (ERP firm / forecast orders; simulated orders live in the grid). Return the row count."""
    for c, h in enumerate(ORDER_HEADERS, start=1):
        ws.cell(1, c, h)
    _style_header(ws, 1, len(ORDER_HEADERS))
    r = 2
    for aid in ids:
        ar = result.articles.get(aid)
        if not ar:
            continue
        for e in ar.events:
            if e.kind != "order" or e.order_type == "PLANNED" or (start and e.date < start):
                continue
            ws.append(
                [
                    aid,
                    ar.article.designation,
                    e.order_type,
                    e.ref,
                    e.supplier_id,
                    e.date,
                    e.qty,
                    e.source,
                    "OUI" if e.late else "",
                    None,
                    None,
                    None,
                    None,
                ]
            )
            ws.cell(r, 6).number_format = DATE_FMT
            for c in (6, 7, 10, 11, 12):
                ws.cell(r, c).fill = INPUT_FILL
                ws.cell(r, c).font = BLUE_FONT
            ws.cell(r, 11).number_format = DATE_FMT
            r += 1
    if live:
        for rr in range(2, r + SPARE_ROWS):
            ws.cell(rr, 13, f'=IF(OR(L{rr}="RECUE",L{rr}="ANNULEE"),0,MAX(0,G{rr}-J{rr}))').number_format = QTY_FMT
            ws.cell(rr, 11).number_format = DATE_FMT
    _widths(ws, (13, 28, 11, 18, 12, 13, 12, 10, 8, 12, 14, 18, 12))
    ws.freeze_panes = "A2"
    return r - 2


ENTRY_HEADERS = [
    "Type (COMMANDE/RECEPTION/AJUSTEMENT/PRODUCTION)",
    "Article ou Programme",
    "Fournisseur",
    "Date",
    "Quantité",
    "Commentaire",
    "Référence commande (réception)",
]


def _entries_template(ws) -> None:
    for c, h in enumerate(ENTRY_HEADERS, start=1):
        ws.cell(1, c, h)
    _style_header(ws, 1, len(ENTRY_HEADERS))
    example = [
        "EXEMPLE",
        "P-00001046",
        "S-000545",
        dt.date(2026, 10, 15),
        1600,
        "exemple (type EXEMPLE = ligne ignorée) : remplacer par COMMANDE / RECEPTION / AJUSTEMENT / PRODUCTION",
        "",
    ]
    for c, v in enumerate(example, start=1):
        ws.cell(2, c, v).font = Font(italic=True, color="888888")
    for rr in range(3, ENTRY_FORMAT_ROWS + 1):
        ws.cell(rr, 4).number_format = DATE_FMT
    _widths(ws, (44, 22, 14, 14, 12, 50, 26))


def alerts_workbook(result: MrpResult, ids: list[str] | None = None) -> bytes:
    wb = Workbook()
    _alerts_sheet(wb.active, result, ids or list(result.articles))
    wb.active.title = "ALERTES"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def orders_workbook(result: MrpResult, ids: list[str] | None = None) -> bytes:
    wb = Workbook()
    _orders_sheet(wb.active, result, ids or list(result.articles))
    wb.active.title = "CARNET_COMMANDES"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# Imports
# =============================================================================
@dataclass
class ParsedPdpLine:
    program_id: str
    week_start: dt.date
    qty: float


def parse_week_label(value: Any) -> dt.date | None:
    """``S11-26`` / ``2028W24`` / ``2026-W11`` / ``W11-2026`` / a date → Monday of the ISO week."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return iso_week_monday(value.date())
    if isinstance(value, dt.date):
        return iso_week_monday(value)
    s = str(value).strip().upper()
    m = re.fullmatch(r"S(\d{1,2})-(\d{2,4})", s)
    if m:
        w, y = int(m.group(1)), int(m.group(2))
        return dt.date.fromisocalendar(y if y > 100 else 2000 + y, w, 1)
    m = re.fullmatch(r"(\d{4})-?W(\d{1,2})", s)
    if m:
        return dt.date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1)
    m = re.fullmatch(r"W(\d{1,2})-(\d{4})", s)
    if m:
        return dt.date.fromisocalendar(int(m.group(2)), int(m.group(1)), 1)
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return iso_week_monday(dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
    return None


def parse_pdp_workbook(
    content: bytes, program_names: dict[str, str], sheet: str | None = None
) -> tuple[list[ParsedPdpLine], list[str]]:
    """Parse a PDP workbook. ``program_names`` maps program name **and** id (upper-cased) → program_id."""
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = (
        wb[sheet]
        if sheet and sheet in wb.sheetnames
        else (wb["SOP - PDP"] if "SOP - PDP" in wb.sheetnames else wb[wb.sheetnames[0]])
    )
    rows = list(ws.iter_rows(values_only=True))
    notes: list[str] = []
    lines: list[ParsedPdpLine] = []
    if not rows:
        return lines, ["classeur vide"]
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    # long layout
    if {"program_id", "qty"} <= set(header) and ("week_start" in header or "iso_week" in header):
        ip, iq = header.index("program_id"), header.index("qty")
        iw = header.index("week_start") if "week_start" in header else header.index("iso_week")
        for r in rows[1:]:
            if not r or r[ip] is None:
                continue
            pid = program_names.get(str(r[ip]).strip().upper())
            monday = parse_week_label(r[iw])
            if not pid or not monday:
                notes.append(f"ligne ignorée : {r[ip]!r} / {r[iw]!r}")
                continue
            lines.append(ParsedPdpLine(pid, monday, float(r[iq] or 0)))
        return lines, notes
    # wide (legacy) layout
    week_cols = {}
    for c, h in enumerate(rows[0]):
        if c == 0:
            continue
        monday = parse_week_label(h)
        if monday:
            week_cols[c] = monday
        elif h not in (None, ""):
            notes.append(f"colonne {c + 1} ignorée : {h!r} n'est pas une semaine")
    if not week_cols:
        return lines, ["aucune colonne semaine reconnue (attendu : S11-26, 2026-W11, 2028W24 ou une date)"]
    for r in rows[1:]:
        if not r or r[0] in (None, ""):
            continue
        pid = program_names.get(str(r[0]).strip().upper())
        if not pid:
            notes.append(f"programme inconnu ignoré : {r[0]!r}")
            continue
        for c, monday in week_cols.items():
            v = r[c] if c < len(r) else None
            if v in (None, ""):
                continue
            try:
                lines.append(ParsedPdpLine(pid, monday, float(v)))
            except (TypeError, ValueError):
                notes.append(f"valeur non numérique ignorée : {r[0]} / {monday} = {v!r}")
    return lines, notes


@dataclass
class ParsedEntry:
    kind: str  # COMMANDE | RECEPTION | AJUSTEMENT | PRODUCTION | COMMANDE_SIMULEE
    key: str  # article_id or program_id
    supplier_id: str | None
    date: dt.date
    qty: float
    comment: str
    order_id: str | None = None  # receipt posted against an order (app id or ERP reference)
    proposal_id: str | None = None  # order created from a proposal decision


def _to_date(v: Any) -> dt.date | None:
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        return None


def _to_float(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _text(v: Any) -> str:
    return str(v).strip() if v is not None else ""


def parse_entries_workbook(content: bytes) -> tuple[list[ParsedEntry], list[str]]:
    """Read every planner input of an exported workbook (or of a bare ``SAISIES`` sheet).

    * ``SAISIES`` rows: type, article / program, supplier, date, qty, comment, order reference;
    * ``CARNET_COMMANDES`` rows with a received quantity → receipts against the order reference;
    * ``SIMULATION`` grid, row *Commandes simulées*: one ``COMMANDE_SIMULEE`` entry per period
      (value or empty → the cell of that day is set / cleared; a week maps to its first day).
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    entries: list[ParsedEntry] = []
    notes: list[str] = []

    ws = wb["SAISIES"] if "SAISIES" in wb.sheetnames else wb[wb.sheetnames[0]]
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or r[0] in (None, ""):
            continue
        kind = _text(r[0]).upper()
        if kind == "EXEMPLE":
            continue
        if kind not in ("COMMANDE", "RECEPTION", "AJUSTEMENT", "PRODUCTION"):
            notes.append(f"SAISIES ligne {n} : type inconnu {r[0]!r}")
            continue
        date, qty = _to_date(r[3] if len(r) > 3 else None), _to_float(r[4] if len(r) > 4 else None)
        if date is None or qty is None:
            notes.append(f"SAISIES ligne {n} : date ou quantité invalide")
            continue
        key = _text(r[1] if len(r) > 1 else None)
        if not key:
            notes.append(f"SAISIES ligne {n} : article / programme manquant")
            continue
        entries.append(
            ParsedEntry(
                kind,
                key,
                _text(r[2] if len(r) > 2 else None) or None,
                date,
                qty,
                _text(r[5] if len(r) > 5 else None),
                order_id=_text(r[6] if len(r) > 6 else None) or None,
            )
        )

    if "CARNET_COMMANDES" in wb.sheetnames:
        # columns: A article, D reference, E supplier, F expected date, J received qty, K receipt date
        for n, r in enumerate(wb["CARNET_COMMANDES"].iter_rows(min_row=2, values_only=True), start=2):
            if not r or r[0] in (None, "") or len(r) < 10:
                continue
            qty = _to_float(r[9])
            if not qty or qty <= 0:
                continue
            date = _to_date(r[10] if len(r) > 10 else None) or _to_date(r[5])
            if date is None:
                notes.append(f"CARNET_COMMANDES ligne {n} : date de réception invalide")
                continue
            entries.append(
                ParsedEntry(
                    "RECEPTION",
                    _text(r[0]),
                    _text(r[4]) or None,
                    date,
                    qty,
                    f"réception saisie dans le carnet ({_text(r[3])})",
                    order_id=_text(r[3]) or None,
                )
            )

    if "SIMULATION" in wb.sheetnames:
        rows = list(wb["SIMULATION"].iter_rows(values_only=True))
        if len(rows) > SIM_HEADER_ROWS:
            starts = [_to_date(v) for v in rows[1][SIM_FIRST_COL - 1 :]]
            label = BLOCK[ROW["sim_orders"]][1]
            for r in rows[SIM_HEADER_ROWS:]:
                if not r or len(r) < 3 or _text(r[2]) != label or r[0] in (None, ""):
                    continue
                for j, day in enumerate(starts):
                    if day is None:
                        continue
                    v = r[SIM_FIRST_COL - 1 + j] if len(r) > SIM_FIRST_COL - 1 + j else None
                    qty = _to_float(v)
                    if qty is None and v not in (None, ""):
                        notes.append(
                            f"SIMULATION {r[0]} {day} : valeur non numérique ignorée ({v!r}) – recalculer le classeur"
                        )
                        continue
                    entries.append(
                        ParsedEntry(
                            "COMMANDE_SIMULEE",
                            _text(r[0]),
                            None,
                            day,
                            qty or 0.0,
                            "réimport de la ligne Commandes simulées",
                        )
                    )
    return entries, notes


# V2 supersedes the legacy simulation workbook; separate PDP/orders/alerts exports remain.


def simulation_workbook(result, article_ids=None, granularity="day", start=None, meta=None):
    return export_grid(result, article_ids, granularity, start, meta)[0]
