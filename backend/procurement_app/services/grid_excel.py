"""Daily editable grid and server-validated round-trip; weekly sheets are summaries."""

from __future__ import annotations

import datetime as dt
import io
import math
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter as letter
from openpyxl.workbook.properties import CalcProperties

from ..engine.calendar import WorkCalendar
from .expression import evaluate

INPUTS = {
    "demand": ("Besoin composants", "demand"),
    "firm_flow": ("Commandes fermes — solde à livrer", "supply_firm"),
    "forecast_flow": ("Commandes prévisionnelles — solde à livrer", "supply_forecast"),
    "planned_flow": ("Commandes simulées", "supply_planned"),
    "receipt_flow": ("Réceptions", "receipts"),
    "adjustment_flow": ("Ajustements (+ / −)", "adjustments"),
}
LABELS = [
    *[(k, v[0]) for k, v in INPUTS.items()],
    ("net_firm", "Solde net ferme"),
    ("net_forecast", "Solde net prévisionnel"),
    ("net_sim", "Solde net simulé"),
    ("physical_firm", "Stock physique ferme"),
    ("physical_forecast", "Stock physique prévisionnel"),
    ("physical_sim", "Stock physique simulé"),
    ("shortage_firm", "Manque ferme"),
    ("shortage_forecast", "Manque prévisionnel"),
    ("shortage_sim", "Manque simulé"),
    ("target", "Stock cible"),
    ("coverage", "Couverture simulée (jours)"),
    ("cumulative", "Besoin cumulé"),
]


def checked_workbook(content: bytes):
    if len(content) > 12 * 1024**2:
        raise ValueError("Classeur limité à 12 Mo")
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if len(z.infolist()) > 200 or sum(x.file_size for x in z.infolist()) > 80 * 1024**2:
            raise ValueError("Classeur décompressé trop volumineux")
        if any("vbaProject" in x.filename or "externalLinks/" in x.filename for x in z.infolist()):
            raise ValueError("Macros et liens externes interdits")
    wb = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    if sum(s.max_row * s.max_column for s in wb) > 600000:
        raise ValueError("Classeur trop volumineux")
    return wb


def export_grid(result, article_ids=None, granularity="day", start=None, meta=None, export_id=""):
    ids = [a for a in (article_ids or result.articles) if a in result.articles]
    start = max(start or result.as_of, result.as_of)
    wb = Workbook()
    ws = wb.active
    ws.title = "SIMULATION"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    dates = [d for d in next(iter(result.articles.values())).dates if d >= start] if ids else []
    visible_dates = list(dates)
    dates += result.tail_dates
    if len(ids) * len(dates) * len(LABELS) > 450000:
        raise ValueError("Export limité à 450 000 cellules ; réduire le périmètre")
    if not dates:
        raise ValueError("Aucune période à exporter")
    ws.append(["PROCUREMENT", "Simulation modifiable", "Saisir dans les cellules bleues", "Initial"])
    ws.append(["Article", "Désignation · unité", "Flux / indicateur", "Veille du début", *dates])
    cal = WorkCalendar.from_spec(result.params.working_weekdays, (meta or {}).get("holidays", []))
    ws.append(["", "", "Jour ouvert", "", *[int(cal.is_working_day(d)) for d in dates]])
    ws.row_dimensions[3].hidden = True
    manifest = {
        "format": 2,
        "export_id": export_id,
        "inputs": {},
        "dates": [d.isoformat() for d in visible_dates],
        "rows": {},
    }
    for aid in ids:
        ar = result.articles[aid]
        a = ar.article
        top = ws.max_row + 2
        rows = {k: top + j for j, (k, _) in enumerate(LABELS)}
        manifest["rows"][aid] = rows
        for k, label in LABELS:
            r = rows[k]
            ws.cell(r, 1, aid).data_type = "s"
            ws.cell(r, 2, f"{a.designation} · {a.unit}").data_type = "s"
            ws.cell(r, 3, label)
        first_idx = ar.dates.index(dates[0])
        for layer, key in [("firm", "stock_firm_net"), ("forecast", "stock_forecast_net"), ("sim", "stock_sim_net")]:
            ws.cell(rows["net_" + layer], 4, getattr(ar, key)[first_idx - 1] if first_idx else ar.kpis["stock_on_hand"])
        ws.cell(rows["target"], 4, a.coverage_target_days)
        ws.row_dimensions[rows["cumulative"]].hidden = True
        for j, d in enumerate(dates, 5):
            col = letter(j)
            prev = letter(j - 1)
            i = ar.dates.index(d) if d in ar.dates else None

            def ref(k):
                return f"{col}{rows[k]}"

            for k, (_, series) in INPUTS.items():
                q = (
                    float(getattr(ar, series)[i])
                    if i is not None
                    else (result.tail_demand[aid][result.tail_dates.index(d)] if k == "demand" else 0.0)
                )
                if i is not None and k == "planned_flow" and result.params.include_proposals_in_simulation:
                    q += ar.supply_proposed[i]
                c = ws.cell(rows[k], j, q)
                c.font = Font(name="Aptos", color="176B88")
                c.fill = PatternFill("solid", fgColor="EAF7FC")
                c.protection = Protection(locked=False)
                if d in visible_dates:
                    manifest["inputs"][c.coordinate] = {"article_id": aid, "date": d.isoformat(), "kind": k, "value": q}
                else:
                    c.protection = Protection(locked=True)
            common = f"{ref('receipt_flow')}+{ref('adjustment_flow')}-{ref('demand')}"
            flows = ref("firm_flow")
            for layer in ("firm", "forecast", "sim"):
                if layer == "forecast":
                    flows += f"+{ref('forecast_flow')}"
                if layer == "sim":
                    flows += f"+{ref('planned_flow')}"
                x = f"{prev}{rows['net_' + layer]}+{flows}+{common}"
                ws.cell(rows["net_" + layer], j, f"=MAX(0,{x})" if result.params.shortage_policy == "lost" else f"={x}")
                ws.cell(rows["physical_" + layer], j, f"=MAX(0,{ref('net_' + layer)})")
                ws.cell(rows["shortage_" + layer], j, f"=MAX(0,-({x}))")
            target_end = (
                cal.add_working_days(d, a.coverage_target_days)
                if result.params.coverage_unit == "working"
                else d + dt.timedelta(days=a.coverage_target_days)
            )
            last = min(len(dates) + 4, j + (target_end - d).days)
            demand_sum = f"SUM({letter(j + 1)}{rows['demand']}:{letter(last)}{rows['demand']})" if last > j else "0"
            target = (
                str(a.safety_stock_qty)
                if result.params.target_policy == "safety_qty"
                else demand_sum
                if result.params.target_policy == "coverage_days"
                else f"{a.safety_stock_qty}+{demand_sum}"
                if result.params.target_policy == "sum"
                else f"MAX({a.safety_stock_qty},{demand_sum})"
            )
            ws.cell(rows["target"], j, "=" + target)
            ws.cell(
                rows["cumulative"], j, f"={ref('demand')}" if j == 5 else f"={prev}{rows['cumulative']}+{ref('demand')}"
            )
            op = "<=" if result.params.coverage_tie_rule == "covered" else "<"
            end = letter(len(dates) + 4)
            if j < len(dates) + 4:
                criteria = f'"{op}"&({ref("cumulative")}+{ref("net_sim")})'
                count = f"COUNTIF({letter(j + 1)}{rows['cumulative']}:{end}{rows['cumulative']},{criteria})"
                if result.params.coverage_unit == "working":
                    count = f"SUMIF({letter(j + 1)}{rows['cumulative']}:{end}{rows['cumulative']},{criteria},{letter(j + 1)}$3:{end}$3)"
                ws.cell(rows["coverage"], j, f"=IF({ref('net_sim')}<0,0,{count})")
            else:
                ws.cell(rows["coverage"], j, 0)
        for row in ws.iter_rows(min_row=top, max_row=top + len(LABELS) - 1, min_col=4, max_col=len(dates) + 4):
            for c in row:
                c.number_format = "#,##0.###;[Red](#,##0.###);–"
        for k in ["physical_firm", "physical_forecast", "physical_sim"]:
            for c in ws[rows[k]]:
                c.font = Font(name="Aptos", bold=True, color="155E63")
                c.fill = PatternFill("solid", fgColor="E7F3EF")
        ws.conditional_formatting.add(
            f"E{rows['shortage_firm']}:{letter(len(dates) + 4)}{rows['shortage_sim']}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FDE4DF")),
        )
    ws.freeze_panes = "E4"
    ws.sheet_view.showGridLines = False
    for c, w in [("A", 18), ("B", 34), ("C", 44), ("D", 16)]:
        ws.column_dimensions[c].width = w
    for j in range(5, len(dates) + 5):
        ws.column_dimensions[letter(j)].width = 13
        ws.cell(2, j).number_format = "ddd dd/mm"
    for row in (1, 2):
        ws.row_dimensions[row].height = 28
        for c in ws[row]:
            c.fill = PatternFill("solid", fgColor="143D48")
            c.font = Font(name="Aptos", bold=True, color="FFFFFF")
    for j in range(5 + len(visible_dates), 5 + len(dates)):
        ws.column_dimensions[letter(j)].hidden = True
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    guide = wb.create_sheet("LIRE_MOI")
    for line in [
        "PROCUREMENT • Simulation Excel",
        "Toutes les saisies de flux se font directement dans SIMULATION (cellules bleues).",
        "Les commandes représentent les SOLDES restant à livrer : lors d’une réception, diminuer le solde et saisir la réception à sa date.",
        "Réceptions et commandes sont des flux distincts : ne jamais conserver la même quantité dans les deux.",
        "Les commandes simulées et ajustements sont signés ; besoins, soldes fermes, prévisions et réceptions sont positifs.",
        "Une expression arithmétique simple (=2*600-50) est acceptée ; les références à d’autres cellules ne sont pas réimportables.",
        "Les colonnes masquées prolongent le PDP pour calculer correctement les cibles en fin d’horizon.",
        "Les calculs restent quotidiens. HEBDOMADAIRE présente les sommes des flux et les stocks de fin de semaine.",
        "Couverture = jours futurs couverts dans le PDP disponible ; elle est bornée par la fin de fenêtre, ce n’est pas une garantie au-delà.",
        f"Politique de manque : {result.params.shortage_policy} ; cible : {result.params.target_policy} ; couverture : {result.params.coverage_unit}.",
        "Réimporter crée un scénario séparé et conserve la base. Aucun ordre ERP ni réception réelle n’est publié.",
        "Les formules se recalculent à l’ouverture dans Excel. Le serveur relance son moteur après import.",
        "Les règles, dates, références et cellules de calcul sont protégées contre les modifications accidentelles.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 130
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical="center")
        guide.row_dimensions[row[0].row].height = 34
    if granularity == "week":
        week = wb.create_sheet("HEBDOMADAIRE", 1)
        week.append(["Article", "Indicateur", "Unité"])
        groups = {}
        for j, d in enumerate(visible_dates, 5):
            groups.setdefault(f"{d.isocalendar().year}-S{d.isocalendar().week:02}", []).append(j)
        for j, label in enumerate(groups, 4):
            week.cell(1, j, label)
        for aid, rows in manifest["rows"].items():
            for k, label in LABELS[:-1]:
                rr = week.max_row + 1
                week.cell(rr, 1, aid).data_type = "s"
                week.cell(rr, 2, label)
                week.cell(rr, 3, "jours" if k == "coverage" else result.articles[aid].article.unit).data_type = "s"
                for c, cols in enumerate(groups.values(), 4):
                    rng = f"SIMULATION!{letter(cols[0])}{rows[k]}:{letter(cols[-1])}{rows[k]}"
                    additive = k in INPUTS or (k.startswith("shortage") and result.params.shortage_policy == "lost")
                    week.cell(rr, c, f"=SUM({rng})" if additive else f"=SIMULATION!{letter(cols[-1])}{rows[k]}")
        week.freeze_panes = "D2"
        week.column_dimensions["A"].width = 18
        week.column_dimensions["B"].width = 44
    fmt = wb.create_sheet("_FORMAT")
    fmt.append(["PROCUREMENT_GRID", 2, export_id])
    fmt.sheet_state = "veryHidden"
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue(), manifest


def parse_grid(content, manifest):
    wb = checked_workbook(content)
    if "_FORMAT" not in wb or list(wb["_FORMAT"].values)[0][:3] != ("PROCUREMENT_GRID", 2, manifest["export_id"]):
        raise ValueError("Classeur non reconnu ; exporter à nouveau depuis PROCUREMENT")
    ws = wb["SIMULATION"]
    for j, d in enumerate(manifest["dates"], 5):
        val = ws.cell(2, j).value
        if not isinstance(val, (dt.date, dt.datetime)) or val.isoformat()[:10] != d:
            raise ValueError("Dates de simulation modifiées")
    changes = []
    for coord, info in manifest["inputs"].items():
        cell = ws[coord]
        if ws.cell(cell.row, 1).value != info["article_id"]:
            raise ValueError("Références articles modifiées")
        raw = cell.value
        if raw in ("", None):
            qty = 0.0
        elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
            qty = float(raw)
        elif isinstance(raw, str):
            qty = evaluate(raw.lstrip("="))
        else:
            raise ValueError(f"Quantité invalide {coord}")
        if not math.isfinite(qty) or abs(qty) > 1e12:
            raise ValueError(f"Quantité excessive {coord}")
        if info["kind"] not in ("planned_flow", "adjustment_flow") and qty < 0:
            raise ValueError(f"Quantité négative interdite {coord}")
        if not math.isclose(qty, info["value"], rel_tol=0, abs_tol=1e-8):
            changes.append(info | {"qty": qty, "cell": coord})
    return changes
