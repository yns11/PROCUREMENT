from io import BytesIO
import openpyxl
import pytest
from procurement.models import Dataset
from procurement.engine import calculate
from procurement.exchange import export_workbook,read_workbook,read_pdp

def test_roundtrip_and_formula_injection(raw):
    raw['items'][0]['name']='=HYPERLINK("https://example.com")'
    data=Dataset.model_validate(raw);content=export_workbook(data,calculate(data))
    assert read_workbook(content)==data
    w=openpyxl.load_workbook(BytesIO(content));assert w['inputs_items']['B2'].data_type=='s'

def test_reject_formulas(raw):
    d=Dataset.model_validate(raw);w=openpyxl.load_workbook(BytesIO(export_workbook(d,calculate(d))))
    w['inputs_plans']['C2']='=1+2';out=BytesIO();w.save(out)
    with pytest.raises(ValueError,match='Formule'):read_workbook(out.getvalue())

def test_csv_iso_monday():
    assert read_pdp(b'program,week,quantity\nM,2026-01-05,0\n')[0]['quantity']=='0'
    with pytest.raises(ValueError):read_pdp(b'program,week,quantity\nM,2026-01-06,1\n')

def test_filtered_outputs_keep_complete_inputs(raw):
    d=Dataset.model_validate(raw);content=export_workbook(d,calculate(d),set(),'month')
    assert read_workbook(content)==d
    w=openpyxl.load_workbook(BytesIO(content));assert w['outputs_series'].max_row==1
