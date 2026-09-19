"""Read-only extraction of legacy workbook inputs and unresolved migration fields.
Usage: python -m scripts.inspect_legacy input.xlsx --output /private/staging.json
Never commit the generated file to a public repository.
"""
import argparse
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
import openpyxl

def inspect(path):
    wb=openpyxl.load_workbook(path,data_only=False)
    cached=openpyxl.load_workbook(path,data_only=True)
    required={'BASE ARTICLE','BOM','SOP - PDP','PREVU - ENGAGÉ','SIMULATION'}
    if not required.issubset(wb.sheetnames):raise ValueError('Structure du classeur non reconnue')
    report={'schema':'legacy-staging/1','publishable':False,'sheets':[],'broken_names':[], 'base_article':[], 'bom':[], 'weekly_plan':[], 'migration_blockers':['Valider le sens de ENGAGÉ avant mapping vers production réelle','Fournir les délais et multiples fournisseurs','Rapprocher stock initial, carnet ouvert et réceptions avec ERP','Résoudre les unités et les multisources ; aucune sélection automatique','Les cellules PDP vides restent absentes et ne sont pas converties en zéro']}
    for ws in wb:
        formulas=sum(c.data_type=='f' for row in ws for c in row)
        errors=sum(c.data_type=='e' for row in cached[ws.title] for c in row)
        report['sheets'].append({'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,'formulas':formulas,'cached_errors':errors})
    report['broken_names']=[{'name':k,'reference':v.attr_text} for k,v in wb.defined_names.items() if '#REF!' in v.attr_text]
    ws=cached['BASE ARTICLE'];headers=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[0]:report['base_article'].append(dict(zip(headers,row)))
    for row in cached['BOM'].iter_rows(min_row=2,max_col=5,values_only=True):
        if row[0]:report['bom'].append(dict(zip(['program','parent','component','quantity','unit'],row)))
    ws=cached['SOP - PDP'];headers=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]:continue
        for week,q in zip(headers[1:],row[1:]):
            if q is not None:report['weekly_plan'].append({'program':row[0],'legacy_week':week,'quantity':q})
    report['multi_source_items']=[k for k,v in Counter(r['REF'] for r in report['base_article']).items() if v>1]
    wb.close();cached.close();return report

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('input');p.add_argument('--output',required=True);a=p.parse_args()
    Path(a.output).write_text(json.dumps(inspect(a.input),ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print('Extraction locale terminée. Résoudre les migration_blockers avant reprise.')
